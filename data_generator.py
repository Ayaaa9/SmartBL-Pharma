from __future__ import annotations

import csv
import hashlib
import random
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional

from faker import Faker
from openpyxl import load_workbook

from config import DatasetConfig, DatasetV2Config
from supplier_template_specs import (
    layout_variant_for_supplier,
    supplier_profile_for_supplier,
    template_family_for_supplier,
    template_layout_for_supplier,
    supplier_template_path,
    supplier_category,
)

TVA_RATES = [0.0, 0.07, 0.10, 0.20]
TEMPLATE_KEYS: List[str] = ["ledger", "clinical", "receipt", "invoice", "dossier", "executive", "compact", "technical", "premium"]
SUPPLIER_TEMPLATE_HINTS: Dict[str, str] = {
    "COOPER": "ledger",
    "DIPHARM": "clinical",
    "REPHAL": "receipt",
    "SARP": "invoice",
    "S.A.R.P": "invoice",
    "KENIPHARMA": "compact",
    "PHARMANORD": "technical",
    "BUMAGECO": "executive",
}


@dataclass
class MedicineRecord:
    normalized_name: str
    base_price: float
    price_min: float
    price_max: float


@dataclass
class BLDataGenerator:
    cfg: DatasetConfig | DatasetV2Config

    def __post_init__(self) -> None:
        self.rng = random.Random(self.cfg.random_seed)
        self.faker = Faker("fr_FR")
        self.faker.seed_instance(self.cfg.random_seed)

        self.suppliers = self._load_suppliers()
        self.medicines = self._load_medicines()

        if not self.suppliers:
            raise RuntimeError("No suppliers loaded from CSV.")
        if not self.medicines:
            raise RuntimeError("No medicines loaded from Excel/CSV source.")

    def generate_document(self, idx: int, split: str = "train") -> Dict[str, Any]:
        supplier = self.rng.choice(self.suppliers)
        template_file = supplier["template_file"]
        supplier_profile = supplier["supplier_profile"]
        template_family = template_family_for_supplier(supplier["name"])
        template_layout = template_layout_for_supplier(supplier["name"])

        issue_date = self._generate_date()
        bl_number = self._generate_bl_number(supplier["name"], idx)
        
        # Vary product count based on supplier category
        category = supplier_profile.get("supplier_category", "pharma_distrib")
        n_items = self._products_count_for_category(category)

        # BL-level lot/expiry patterns
        presence_rate_lot = getattr(self.cfg, "presence_rate_lot", 0.60)
        presence_rate_expiry = getattr(self.cfg, "presence_rate_expiry", 0.65)

        # Determine BL-level modes: some BLs have all lots, none, or mixed
        r = self.rng.random()
        if r < 0.12:
            lot_mode = "all"
        elif r < 0.20:
            lot_mode = "none"
        else:
            lot_mode = "mixed"

        s = self.rng.random()
        if s < 0.12:
            expiry_mode = "all"
        elif s < 0.20:
            expiry_mode = "none"
        else:
            expiry_mode = "mixed"

        hard_case = self.rng.random() < getattr(self.cfg, "hard_case_rate", 0.0)

        products = self._generate_products_for_category(
            n_items,
            supplier_profile,
            category,
            lot_mode=lot_mode,
            expiry_mode=expiry_mode,
            presence_rate_lot=presence_rate_lot,
            presence_rate_expiry=presence_rate_expiry,
            hard_case=hard_case,
        )
        totals = self._compute_totals(products)

        canvas = self.rng.choice(self.cfg.canvas_sizes)
        doc_id = f"bl_{idx:06d}"

        render_context = {
            "doc_id": doc_id,
            "supplier": supplier,
            "supplier_profile": supplier_profile,
            "bl_number": bl_number,
            "issue_date": issue_date,
            "client_name": self._generate_client_name(),
            "client_code": f"C{self.rng.randint(1000, 9999)}",
            "client_address": self.faker.address().replace("\n", ", "),
            "products": products,
            "totals": totals,
            "currency": self.cfg.currency,
            "font_family": self.rng.choice(self.cfg.font_families),
            "header_offset": self.rng.randint(-8, 10),
            "table_skew": round(self.rng.uniform(-0.4, 0.45), 3),
            "show_stamp": self.rng.random() < self.cfg.stamp_probability,
            "show_signature": self.rng.random() < self.cfg.signature_probability,
            "canvas_width": canvas["width"],
            "canvas_height": canvas["height"],
            "column_order": supplier_profile["column_order"],
            "template_family": template_family,
            "template_layout": template_layout,
            "difficulty": self._choose_difficulty(n_items),
            "hard_case": hard_case,
            "product_count": len(products),
            "include_columns": {
                "code_produit": supplier_profile["show_code"],
                "pph": supplier_profile["show_pph"],
                "tva": supplier_profile["show_tva"],
            },
            "category_badge": supplier_profile.get("category_label", ""),
        }
        # Build Wellys-aligned annotation (no totals, no line_total)
        items = []
        for p in products:
            item = {
                "product_name": p.get("raw_text"),
                "quantity": p.get("quantity"),
                "unit_price": p.get("unit_price") or p.get("price"),
                "lot_number": p.get("lot") if p.get("lot") is not None else None,
                "expiry_date": p.get("peremption") if p.get("peremption") is not None else None,
            }
            items.append(item)

        annotation = {
            "bl_number": bl_number,
            "bl_date": issue_date,
            "supplier_name": supplier["name"],
            "supplier_category": category,
            "items": items,
            "difficulty": render_context["difficulty"],
            "hard_case": hard_case,
            "template_family": template_family,
            "template": template_file,
            "supplier_profile": supplier_profile,
            "split": split,
            "review_flags": [],
        }

        return {
            "doc_id": doc_id,
            "template": template_file,
            "render_context": render_context,
            "annotation": annotation,
            "split": split,
        }

    def _load_suppliers(self) -> List[Dict[str, Any]]:
        csv_path = self.cfg.suppliers_csv_path
        if not csv_path.exists():
            fallback = self._first_existing_path(self.cfg.suppliers_csv_fallbacks)
            if fallback is not None:
                csv_path = fallback

        if not csv_path.exists():
            return []

        suppliers: List[Dict[str, Any]] = []
        with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            for row in reader:
                name = (row.get("name") or "").strip().strip(",")
                if not name:
                    continue

                normalized = self._slugify(name).replace("-", "")
                if normalized in {"nonrenseigne", "nonrenseign"}:
                    continue

                template_slug = self._slugify(name)
                template_file = supplier_template_path(name)
                supplier_profile = supplier_profile_for_supplier(name)
                suppliers.append(
                    {
                        "id": int((row.get("id") or "0") or 0),
                        "name": name,
                        "city": (row.get("city") or "").strip() or self.faker.city().upper(),
                        "address": (row.get("address") or "").strip() or self.faker.address().replace("\n", ", "),
                        "phone": self._fake_phone(),
                        "website": (row.get("website") or "").strip(),
                        "responsable": (row.get("pharmacist_responsible") or "").strip(),
                        "template_key": template_slug,
                        "template_file": template_file,
                        "supplier_profile": supplier_profile,
                    }
                )

        return suppliers

    def _load_medicines(self) -> List[MedicineRecord]:
        path = self._first_existing_path(self.cfg.medicines_file_candidates)
        if path is None:
            return self._fallback_medicines()

        if path.suffix.lower() in {".xlsx", ".xls"}:
            rows = self._read_excel_rows(path)
        else:
            rows = self._read_csv_rows(path)

        if not rows:
            return self._fallback_medicines()

        col_map = {c.lower().strip(): c for c in rows[0].keys()}
        name_col = self._pick_existing_col(
            col_map,
            ["name", "designation", "medicament", "medicine", "produit", "libelle"],
        )
        if name_col is None:
            return self._fallback_medicines()

        price_cols = [
            self._pick_existing_col(col_map, ["pph"]),
            self._pick_existing_col(col_map, ["pfht"]),
            self._pick_existing_col(col_map, ["ppv"]),
            self._pick_existing_col(col_map, ["price", "prix", "unit_price", "pu"]),
        ]
        price_cols = [c for c in price_cols if c is not None]

        records: List[MedicineRecord] = []
        seen = set()

        for row in rows:
            raw = str(row.get(name_col, "") or "").strip()
            if not raw or raw.lower() == "nan":
                continue

            normalized_name = self._normalize_name(raw)
            if not normalized_name or normalized_name in seen:
                continue

            base_price = self._extract_price(row, price_cols)
            if base_price is None or base_price <= 0:
                base_price = self.rng.uniform(8.0, 260.0)

            pmin, pmax = self._price_range_from_base(base_price)
            records.append(
                MedicineRecord(
                    normalized_name=normalized_name,
                    base_price=round(base_price, 2),
                    price_min=pmin,
                    price_max=pmax,
                )
            )
            seen.add(normalized_name)

        return records if records else self._fallback_medicines()

    def _generate_products(self, n_items: int, supplier_profile: Optional[Dict[str, Any]] = None, lot_mode: str = "mixed", expiry_mode: str = "mixed", presence_rate_lot: float = 0.6, presence_rate_expiry: float = 0.65) -> List[Dict[str, Any]]:
        selected = self._sample_medicines(n_items, supplier_profile)
        products: List[Dict[str, Any]] = []

        for med in selected:
            quantity = self.generate_realistic_quantity()
            price = self._price_from_medicine(med, category=None)
            tva = self.rng.choice(TVA_RATES)
            remise = round(self.rng.uniform(0.0, 12.0), 1) if self.rng.random() < 0.6 else None
            factor = 1.0 - ((remise or 0.0) / 100.0)
            line_total = round(quantity * price * factor, 2)

            raw_text = self.generate_product_name_variant(med.normalized_name)
            code = f"{self.rng.randint(100000, 999999)}"
            # Lot and expiry presence per BL-level modes and per-item rates
            lot = None
            peremption = None
            if lot_mode == "all":
                lot = self._random_lot()
            elif lot_mode == "none":
                lot = None
            else:
                if self.rng.random() < presence_rate_lot:
                    lot = self._random_lot()

            if expiry_mode == "all":
                peremption = self._random_expiry()
            elif expiry_mode == "none":
                peremption = None
            else:
                if self.rng.random() < presence_rate_expiry:
                    peremption = self._random_expiry()

            product = {
                "raw_text": raw_text,
                "normalized_name": med.normalized_name,
                "designation": raw_text,
                "quantity": quantity,
                "qty": quantity,
                "price": price,
                "unit_price": price,
                "pph": round(price * self.rng.uniform(1.08, 1.22), 2),
                "tva": tva,
                "remise": remise,
                # line_total intentionally included for render only; may be stripped later
                "line_total": line_total,
                "code_produit": code,
                "peremption": peremption,
                "lot": lot,
                "offset_px": self.rng.randint(-4, 4) if self.rng.random() < self.cfg.misalign_probability else 0,
                "is_merged_row": False,
            }

            if self.rng.random() < self.cfg.merged_row_probability:
                alt = self.rng.choice(self.medicines).normalized_name
                product["designation"] = f"PACK {raw_text} + {self._noisy_medicine_text(alt)}"
                product["raw_text"] = product["designation"]
                product["is_merged_row"] = True

            self._maybe_drop_fields(product)
            products.append(product)

        return products

    def _maybe_drop_fields(self, product: Dict[str, Any]) -> None:
        protected = isinstance(self.cfg, DatasetV2Config)
        drop_candidates = ["line_total", "code_produit", "remise"]
        if not protected:
            drop_candidates = ["qty", "unit_price"] + drop_candidates

        for key in drop_candidates:
            if self.rng.random() < self.cfg.missing_cell_probability:
                product[key] = None
                if key == "qty":
                    product["quantity"] = None
                if key == "unit_price":
                    product["price"] = None

    def _compose_product_name(
        self,
        normalized_name: str,
        family: str,
        product_type: str,
        category: str,
    ) -> str:
        descriptor = ""
        if product_type == "parapharmacie":
            descriptor = self.rng.choice(["gel", "crème", "lotion", "spray", "vaporisateur", "bain"])
        elif product_type == "cosmetique_premium":
            descriptor = self.rng.choice(["Sérum", "Crème", "Masque", "Huile", "Lotion", "Émulsion"])

        if category == "ticket_compact":
            raw = f"{family} {normalized_name[:20].strip()}"
        elif category == "dossier_administratif":
            raw = f"[{family}] {normalized_name}"
        else:
            raw = f"{family} - {normalized_name}"

        if descriptor:
            raw = f"{raw} {descriptor}"
        if self.rng.random() < 0.28:
            raw = f"{raw} {self.rng.choice(['Boîte', 'Flacon', 'Sachet', 'Tube'])} {self.rng.choice(['30 cps', '100 ml', '50 ml', '20 cps', '10 x 5 ml'])}"
        return raw.strip()

    def _apply_ocr_confusion(self, text: str) -> str:
        substitutions = {
            "0": self.rng.choice(["O", "D", "Q"]),
            "1": self.rng.choice(["I", "l", "|"]),
            "2": self.rng.choice(["Z", "R"]),
            "5": self.rng.choice(["S", "s"]),
            "8": self.rng.choice(["B", "&"]),
            "O": self.rng.choice(["0", "Q"]),
            "I": self.rng.choice(["1", "l"]),
            "l": self.rng.choice(["1", "I"]),
            "S": self.rng.choice(["5", "s"]),
        }
        out = []
        for ch in text:
            if ch in substitutions and self.rng.random() < 0.18:
                out.append(substitutions[ch])
            else:
                out.append(ch)
        if self.rng.random() < 0.2:
            out = list(" ".join("".join(out).split()))
        return self._apply_spacing_noise("".join(out))

    def _apply_hard_case_to_product(self, product: Dict[str, Any]) -> None:
        anomaly_type = self.rng.choices(
            ["missing_quantity", "missing_price", "zero_quantity", "wrong_total", "ocr_noise"],
            weights=[0.22, 0.22, 0.18, 0.18, 0.20],
            k=1,
        )[0]

        if anomaly_type == "missing_quantity":
            product["quantity"] = None
            product["qty"] = None
        elif anomaly_type == "missing_price":
            product["unit_price"] = None
            product["price"] = None
        elif anomaly_type == "zero_quantity":
            product["quantity"] = 0
            product["qty"] = 0
        elif anomaly_type == "wrong_total":
            if product.get("line_total") is not None:
                product["line_total"] = round(product["line_total"] * self.rng.uniform(0.55, 1.35), 2)
        elif anomaly_type == "ocr_noise":
            product["designation"] = self._apply_ocr_confusion(product["designation"])
            product["raw_text"] = product["designation"]

        # Introduce occasional challenges in lot/peremption fields too
        if self.rng.random() < 0.35:
            product["lot"] = None if self.rng.random() < 0.6 else f"{self.rng.choice(['ERR', '??', 'N/A'])}{self.rng.randint(10,99)}"
        if self.rng.random() < 0.35:
            product["peremption"] = None if self.rng.random() < 0.6 else self._random_expiry()

    def generate_realistic_quantity(self, product_category: str = "medicament", difficulty: str = "medium") -> Any:
        # Returns an integer or a decimal quantity according to configured rates
        int_rate = getattr(self.cfg, "integer_quantity_rate", 0.85)

        # Difficulty may bias towards higher counts or decimals
        if difficulty == "easy":
            choice_pool = [1, 2, 3, 5, 6]
        elif difficulty == "medium":
            choice_pool = [1, 2, 3, 5, 6, 10, 12]
        else:
            choice_pool = [1, 2, 3, 5, 6, 10, 12, 20, 24, 30, 50]

        if self.rng.random() < int_rate:
            return int(self.rng.choice(choice_pool))

        # decimal quantities: pick from common fractional values
        decimals = [0.5, 0.25, 0.75, 1.5]
        base = int(self.rng.choice(choice_pool))
        return round(base + self.rng.choice(decimals), 2)

    def _random_lot(self) -> str:
        patterns = [
            lambda: f"{self.rng.choice('ABCDEFGHIJKLMNOPQRSTUVWXYZ')}{self.rng.randint(10,999):03d}{self.rng.choice('ABCDEFGHIJKLMNOPQRSTUVWXYZ')}",
            lambda: f"B{self.rng.randint(2000,2029)}-{self.rng.randint(100,999)}",
            lambda: f"LOT-{self.rng.randint(10000,99999)}",
            lambda: f"P{self.rng.randint(1000,9999)}",
            lambda: f"{self.rng.randint(100000,999999)}",
            lambda: f"{self.rng.randint(10000000,99999999)}",
            lambda: f"Y{self.rng.randint(1000,9999)}",
        ]
        return self.rng.choice(patterns)()

    def _random_expiry(self) -> str:
        # prefer MM/YY or MM/YYYY, with DD/MM/YYYY rare
        months = self.rng.randint(getattr(self.cfg, "expiry_min_months", 6), getattr(self.cfg, "expiry_max_months", 36))
        d = date.today() + timedelta(days=months * 30)
        formats = getattr(self.cfg, "expiry_formats", ["%m/%y", "%m/%Y"])
        fmt = self.rng.choices(formats, weights=[0.45, 0.45, 0.10] if len(formats) >= 3 else None, k=1)[0]
        return d.strftime(fmt)

    def _compute_totals(self, products: List[Dict[str, Any]]) -> Dict[str, float]:
        ht = 0.0
        tva_sum = 0.0
        total_qty = 0

        for p in products:
            qty = p.get("quantity")
            price = p.get("price")
            if qty is None or price is None:
                continue
            line_ht = qty * price
            remise = p.get("remise") or 0.0
            line_ht *= (1.0 - remise / 100.0)
            ht += line_ht
            tva_sum += line_ht * float(p.get("tva") or 0.0)
            total_qty += int(qty)

        ttc = ht + tva_sum
        return {
            "ht": round(ht, 2),
            "tva": round(tva_sum, 2),
            "ttc": round(ttc, 2),
            "total_qty": total_qty,
        }

    def _generate_client_name(self) -> str:
        prefix = self.rng.choice(["PHARMACIE", "PHARMACIE CENTRALE", "PARAPHARMACIE"])
        return f"{prefix} {self.faker.last_name().upper()}"

    def _generate_bl_number(self, supplier_name: str, idx: int) -> str:
        code = "".join(c for c in supplier_name if c.isalpha())[:4].upper()
        serial = self.rng.randint(100000, 999999)
        return f"{code}-{date.today().year % 100:02d}-{idx:05d}-{serial}"

    def _generate_date(self) -> str:
        days_back = self.rng.randint(0, 730)
        d = date.today() - timedelta(days=days_back)
        return d.strftime("%d/%m/%Y")

    def _noisy_medicine_text(self, normalized_name: str) -> str:
        text = normalized_name

        if self.rng.random() < 0.7:
            text = self._apply_abbreviation(text)
        if self.rng.random() < 0.6:
            text = self._apply_spacing_noise(text)
        if self.rng.random() < 0.5:
            text = self._remove_accents(text)
        if self.rng.random() < 0.55:
            text = self._apply_case_noise(text)
        if self.rng.random() < 0.45:
            text = self._apply_spelling_noise(text)

        return re.sub(r"\s+", " ", text).strip()

    def _apply_case_noise(self, text: str) -> str:
        mode = self.rng.choice(["upper", "lower", "title", "mixed"])
        if mode == "upper":
            return text.upper()
        if mode == "lower":
            return text.lower()
        if mode == "title":
            return text.title()

        chars = []
        for ch in text:
            if ch.isalpha() and self.rng.random() < 0.35:
                chars.append(ch.lower())
            elif ch.isalpha() and self.rng.random() < 0.35:
                chars.append(ch.upper())
            else:
                chars.append(ch)
        return "".join(chars)

    def _apply_abbreviation(self, text: str) -> str:
        replacements = {
            " MILIGRAMME": " MG",
            " MILLIGRAMME": " MG",
            " MILLIGRAMMES": " MG",
            " MG": self.rng.choice(["MG", "mg", " Mg"]),
            " CPR": self.rng.choice([" CPR", " CP", " CPR."]),
            " COMPRIME": self.rng.choice([" CPR", " CP"]),
            " GRAMME": " G",
            " 1000 MG": self.rng.choice([" 1000MG", " 1G"]),
            " SOLUTION INJECTABLE": self.rng.choice([" SOL. INJ", " INJ", " SOL INJECT"]),
        }

        out = f" {text} "
        for src, dst in replacements.items():
            if src in out:
                out = out.replace(src, dst)
        return out.strip()

    def _apply_spelling_noise(self, text: str) -> str:
        if len(text) < 6:
            return text

        chars = list(text)
        n_err = self.rng.randint(1, 2)
        
        for _ in range(n_err):
            alpha_indices = [i for i, ch in enumerate(chars) if ch.isalpha()]
            if not alpha_indices:
                break
            
            idx = self.rng.choice(alpha_indices)
            op = self.rng.choice(["delete", "swap", "replace"])

            if op == "delete" and len(chars) > 4 and idx < len(chars):
                chars.pop(idx)
            elif op == "swap" and idx < len(chars) - 1 and chars[idx + 1].isalpha():
                chars[idx], chars[idx + 1] = chars[idx + 1], chars[idx]
            elif op == "replace" and idx < len(chars):
                chars[idx] = self.rng.choice("ABCDEFGHIJKLMNOPQRSTUVWXYZ")

        return "".join(chars)

    def _apply_spacing_noise(self, text: str) -> str:
        out = text
        out = re.sub(r"(\d+)\s*MG", lambda m: f"{m.group(1)}{self.rng.choice(['MG', ' mg', '  MG'])}", out)
        out = re.sub(r"\s+", " ", out)

        if self.rng.random() < 0.3 and len(out) > 8:
            pos = self.rng.randint(3, len(out) - 3)
            out = out[:pos] + " " + out[pos:]
        return out

    def _remove_accents(self, text: str) -> str:
        normalized = unicodedata.normalize("NFKD", text)
        return normalized.encode("ascii", "ignore").decode("ascii")

    def _normalize_name(self, text: str) -> str:
        out = self._remove_accents(str(text).strip())
        out = re.sub(r"\s+", " ", out)
        return out.upper()

    def _sample_medicines(self, n_items: int, supplier_profile: Optional[Dict[str, Any]] = None) -> List[MedicineRecord]:
        band = (supplier_profile or {}).get("medicine_band", "balanced")
        if band == "low":
            pool = [m for m in self.medicines if m.base_price < 70] or self.medicines
        elif band == "mid":
            pool = [m for m in self.medicines if 50 <= m.base_price <= 140] or self.medicines
        elif band == "premium":
            pool = [m for m in self.medicines if m.base_price >= 90] or self.medicines
        else:
            pool = self.medicines

        if n_items <= len(pool):
            return self.rng.sample(pool, k=n_items)

        picks = pool.copy()
        while len(picks) < n_items:
            picks.append(self.rng.choice(pool))
        self.rng.shuffle(picks)
        return picks[:n_items]

    def _sample_quantity(self) -> int:
        roll = self.rng.random()
        if roll < 0.65:
            return self.rng.randint(1, 6)
        if roll < 0.9:
            return self.rng.randint(7, 12)
        return self.rng.randint(13, 20)

    def _price_range_from_base(self, base_price: float) -> tuple[float, float]:
        spread = max(1.0, base_price * self.rng.uniform(0.08, 0.25))
        pmin = max(2.0, base_price - spread)
        pmax = base_price + spread
        return (round(pmin, 2), round(pmax, 2))

    def _extract_price(self, row: Dict[str, Any], price_cols: List[str]) -> Optional[float]:
        for col in price_cols:
            val = row.get(col)
            if val is None:
                continue
            number = self._parse_number(val)
            if number is not None and number > 0:
                return number
        return None

    def _parse_number(self, val: Any) -> Optional[float]:
        if isinstance(val, (float, int)):
            return float(val)

        s = str(val).strip().upper()
        s = s.replace("MAD", "").replace("DHS", "").replace(" ", "")
        if s.count(",") == 1 and s.count(".") == 0:
            s = s.replace(",", ".")
        try:
            return float(s)
        except ValueError:
            return None

    def _read_excel_rows(self, path: Path) -> List[Dict[str, Any]]:
        wb = load_workbook(path, data_only=True, read_only=True)
        target = "products_export" if "products_export" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[target]

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        out: List[Dict[str, Any]] = []
        for raw in rows[1:]:
            item = {}
            for i, h in enumerate(headers):
                if not h:
                    continue
                item[h] = raw[i] if i < len(raw) else None
            if item:
                out.append(item)
        return out

    def _read_csv_rows(self, path: Path) -> List[Dict[str, Any]]:
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))

    def _fallback_medicines(self) -> List[MedicineRecord]:
        defaults = [
            ("DOLIPRANE 1000 MG CPR", 41.0),
            ("AUGMENTIN 1 G CPR", 89.0),
            ("VENTOLINE AEROSOL", 78.0),
            ("SPASFON LYOC", 35.0),
            ("SMECTA SACHET", 52.0),
            ("GLUCOPHAGE 850 MG", 64.0),
            ("EUTHYROX 50 MCG", 31.0),
            ("RHINADVIL RHUME", 71.0),
            ("MOTILIUM 10 MG", 58.0),
            ("LIPANTHYL 200 MG", 112.0),
        ]
        out = []
        for name, base in defaults:
            pmin, pmax = self._price_range_from_base(base)
            out.append(MedicineRecord(name, base, pmin, pmax))
        return out

    def _supplier_profile(self, template_key: str) -> Dict[str, Any]:
        base = {
            "pharmanord": {
                "doc_title": "BON DE LIVRAISON",
                "header_color": "#334d7a",
                "column_order": ["code", "designation", "qty", "pph", "unit_price", "line_total"],
                "totals_labels": ["TOTAL HT", "TVA", "NET A PAYER"],
                "footer": "Les produits ne sont ni repris ni echanges apres 48 heures.",
            },
            "bumageco": {
                "doc_title": "BON DE LIVRAISON / FACTURE",
                "header_color": "#385f54",
                "column_order": ["designation", "qty", "unit_price", "tva", "line_total"],
                "totals_labels": ["MONTANT HT", "MONTANT TVA", "TOTAL TTC"],
                "footer": "Reglement conforme aux conditions generales de vente.",
            },
            "kenipharma": {
                "doc_title": "BL - KENIPHARMA",
                "header_color": "#3b3b3b",
                "column_order": ["code", "designation", "qty", "unit_price", "line_total", "peremption"],
                "totals_labels": ["SOUS TOTAL", "TVA", "MONTANT TTC"],
                "footer": "Controle qualite effectue a la livraison.",
            },
            "rephal": {
                "doc_title": "BON DE REPARTITION",
                "header_color": "#604033",
                "column_order": ["designation", "lot", "qty", "remise", "unit_price", "line_total"],
                "totals_labels": ["TOTAL HT", "TVA FACTUREE", "TOTAL TTC"],
                "footer": "Bon conforme au referentiel interne du repartiteur.",
            },
            "sarp": {
                "doc_title": "BON DE LIVRAISON SARP",
                "header_color": "#6c3f66",
                "column_order": ["designation", "qty", "unit_price", "pph", "tva", "line_total", "code"],
                "totals_labels": ["TOTAL H.T", "TVA", "TOTAL GENERAL"],
                "footer": "Merci de verifier les quantites et dates de peremption a reception.",
            },
        }
        return base[template_key]

    def _choose_template_for_supplier(self, supplier_name: str) -> str:
        upper_name = supplier_name.upper()
        for key, template in SUPPLIER_TEMPLATE_HINTS.items():
            if key in upper_name:
                return template
        return self.rng.choice(TEMPLATE_KEYS)

    def _pick_existing_col(self, col_map: Dict[str, str], names: List[str]) -> Optional[str]:
        for n in names:
            if n in col_map:
                return col_map[n]
        return None

    def _first_existing_path(self, paths: List[Path]) -> Optional[Path]:
        for p in paths:
            if p.exists():
                return p
        return None

    def _fake_phone(self) -> str:
        return f"05 {self.rng.randint(20, 39)} {self.rng.randint(10, 99)} {self.rng.randint(10, 99)} {self.rng.randint(10, 99)}"

    def _future_month_year(self) -> str:
        months = self.rng.randint(3, 40)
        d = date.today() + timedelta(days=months * 30)
        return d.strftime("%m/%Y")

    def _slugify(self, txt: str) -> str:
        normalized = unicodedata.normalize("NFKD", txt)
        ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
        ascii_text = ascii_text.lower().strip()
        ascii_text = re.sub(r"[^a-z0-9]+", "-", ascii_text)
        ascii_text = re.sub(r"-{2,}", "-", ascii_text).strip("-")
        return ascii_text or "supplier"

    def _products_count_for_category(self, category: str) -> int:
        """Generate product count based on supplier category."""
        if isinstance(self.cfg, DatasetV2Config):
            roll = self.rng.random()
            max_items = self.cfg.max_items_per_bl
            if roll < self.cfg.distribution_small:
                return self.rng.randint(3, min(6, max_items))
            if roll < self.cfg.distribution_small + self.cfg.distribution_medium:
                return self.rng.randint(7, min(12, max_items))
            if roll < self.cfg.distribution_small + self.cfg.distribution_medium + self.cfg.distribution_large:
                return self.rng.randint(13, min(20, max_items))
            return self.rng.randint(21, min(30, max_items))

        if category == "pharma_distrib":
            return self.rng.randint(8, 20)
        elif category == "facture_premium":
            return self.rng.randint(5, 15)
        elif category == "ticket_compact":
            return self.rng.randint(1, 8)
        elif category == "dossier_administratif":
            return self.rng.randint(3, 12)
        else:
            return self.rng.randint(self.cfg.min_products_per_doc, self.cfg.max_products_per_doc)

    def _generate_products_for_category(
        self,
        n_items: int,
        supplier_profile: Dict[str, Any],
        category: str,
        lot_mode: str = "mixed",
        expiry_mode: str = "mixed",
        presence_rate_lot: float = 0.6,
        presence_rate_expiry: float = 0.65,
        hard_case: bool = False,
    ) -> List[Dict[str, Any]]:
        """Generate products with category-specific prefixes, product types and realistic anomalies."""
        medicines = self._sample_medicines(n_items, supplier_profile)
        families = self.cfg.product_families.get(category, self.cfg.product_families["pharma_distrib"])
        product_mix = getattr(self.cfg, "product_category_mix", {"medicament": 0.7, "parapharmacie": 0.2, "cosmetique_premium": 0.1})

        products: List[Dict[str, Any]] = []
        for medicine in medicines:
            product_type = self.rng.choices(
                population=list(product_mix.keys()),
                weights=list(product_mix.values()),
                k=1,
            )[0]
            family = self.rng.choice(families)
            difficulty = self._choose_difficulty(n_items)
            quantity = self.generate_realistic_quantity(product_category=product_type, difficulty=difficulty)

            raw_text = self._compose_product_name(medicine.normalized_name, family, product_type, category)
            raw_text = self.generate_product_name_variant(raw_text)
            if self.rng.random() < getattr(self.cfg, "ocr_confusion_rate", 0.0):
                raw_text = self._apply_ocr_confusion(raw_text)

            base_price = self._price_from_medicine(medicine, product_type)
            if product_type == "cosmetique_premium":
                variance = self.rng.uniform(1.05, 1.40)
            elif product_type == "parapharmacie":
                variance = self.rng.uniform(0.95, 1.18)
            else:
                variance = self.rng.uniform(0.85, 1.18)

            price = round(max(2.0, base_price * variance), 2)
            if category == "dossier_administratif":
                tva = self.rng.choice([0.0, 0.07, 0.10, 0.20])
            else:
                tva = self.rng.choice([0.07, 0.10, 0.20])

            remise = round(self.rng.uniform(0.0, 12.0), 1) if self.rng.random() < 0.55 else None
            factor = 1.0 - ((remise or 0.0) / 100.0)
            line_total = None
            if quantity is not None and price is not None:
                try:
                    line_total = round(quantity * price * factor, 2)
                except Exception:
                    line_total = None

            code = f"{self.rng.randint(100000, 999999)}"
            lot = None
            peremption = None
            if lot_mode == "all":
                lot = self._random_lot()
            elif lot_mode == "mixed" and self.rng.random() < presence_rate_lot:
                lot = self._random_lot()

            if expiry_mode == "all":
                peremption = self._random_expiry()
            elif expiry_mode == "mixed" and self.rng.random() < presence_rate_expiry:
                peremption = self._random_expiry()

            product = {
                "raw_text": raw_text,
                "normalized_name": medicine.normalized_name,
                "designation": raw_text,
                "quantity": quantity,
                "qty": quantity,
                "price": price,
                "unit_price": price,
                "pph": round(price * self.rng.uniform(1.08, 1.22), 2),
                "tva": tva,
                "remise": remise,
                "line_total": line_total,
                "code_produit": code,
                "peremption": peremption,
                "lot": lot,
                "offset_px": self.rng.randint(-4, 4) if self.rng.random() < self.cfg.misalign_probability else 0,
                "is_merged_row": False,
                "family": family,
                "product_type": product_type,
            }

            if self.rng.random() < self.cfg.merged_row_probability:
                alt = self.rng.choice(self.medicines).normalized_name
                product["designation"] = f"PACK {raw_text} + {self._noisy_medicine_text(alt)}"
                product["raw_text"] = product["designation"]
                product["is_merged_row"] = True

            if hard_case and self.rng.random() < getattr(self.cfg, "error_item_rate", 0.0):
                self._apply_hard_case_to_product(product)

            self._maybe_drop_fields(product)
            products.append(product)

        return products

    def _choose_difficulty(self, n_items: int) -> str:
        if n_items <= 6:
            return self.rng.choices(["easy", "medium"], weights=[0.8, 0.2], k=1)[0]
        if n_items <= 12:
            return self.rng.choices(["medium", "easy", "hard"], weights=[0.7, 0.2, 0.1], k=1)[0]
        return self.rng.choices(["hard", "medium"], weights=[0.8, 0.2], k=1)[0]

    def _price_from_medicine(self, medicine: MedicineRecord, category: Optional[str] = None) -> float:
        # Use medicine.base_price as anchor and apply a controlled variation inside a realistic band
        base = getattr(medicine, "base_price", None) or (medicine.price_min + medicine.price_max) / 2

        band_map = {
            "pharma_distrib": (5.0, 300.0),
            "facture_premium": (20.0, 800.0),
            "ticket_compact": (5.0, 300.0),
            "dossier_administratif": (80.0, 1500.0),
            "parapharmacie": (10.0, 420.0),
            "cosmetique_premium": (20.0, 1200.0),
            "medicament": (5.0, 350.0),
        }
        band_min, band_max = band_map.get(category or "pharma_distrib", (5.0, 300.0))

        if base < band_min:
            base = band_min + self.rng.uniform(0.0, 18.0)
        elif base > band_max:
            base = band_max - self.rng.uniform(0.0, max(2.0, band_max * 0.08))

        adj = self.rng.uniform(0.96, 1.08)
        if category == "facture_premium":
            adj = self.rng.uniform(0.98, 1.12)
        elif category == "ticket_compact":
            adj = self.rng.uniform(0.92, 1.02)
        elif category == "dossier_administratif":
            adj = self.rng.uniform(0.95, 1.10)
        elif category == "parapharmacie":
            adj = self.rng.uniform(0.98, 1.14)
        elif category == "cosmetique_premium":
            adj = self.rng.uniform(1.00, 1.18)

        candidate = round(base * adj, 2)
        pmin = getattr(medicine, "price_min", None)
        pmax = getattr(medicine, "price_max", None)
        if pmin is not None:
            candidate = max(candidate, pmin)
        if pmax is not None:
            candidate = min(candidate, pmax)
        candidate = max(band_min, min(candidate, band_max))
        return round(candidate, 2)

    def generate_product_name_variant(self, original_name: str) -> str:
        # Build variants using existing noise functions but keep this as explicit API
        name = original_name
        # 40% keep exact, otherwise apply one or more transformations
        if self.rng.random() < 0.4:
            return name

        if self.rng.random() < 0.6:
            name = self._apply_abbreviation(name)
        if self.rng.random() < 0.6:
            name = self._apply_spacing_noise(name)
        if self.rng.random() < 0.5:
            name = self._remove_accents(name)
        if self.rng.random() < 0.5:
            name = self._apply_case_noise(name)
        if self.rng.random() < 0.35:
            name = self._apply_spelling_noise(name)

        # Occasionally move dosage or remove units
        if self.rng.random() < 0.12:
            name = re.sub(r"(\d+\s*(MG|G|ML|UG))", "", name)

        return re.sub(r"\s+", " ", name).strip()
